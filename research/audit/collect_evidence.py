from pathlib import Path
import json, csv, hashlib, subprocess, sys, zipfile
import yaml, fitz
from bs4 import BeautifulSoup
ROOT=Path(__file__).resolve().parents[1]
OUT=Path(__file__).resolve().parent
(OUT/'extracted').mkdir(exist_ok=True)
manifest=[]
for p in sorted(ROOT.rglob('*')):
    if p.is_file() and OUT not in p.parents:
        manifest.append({'file':str(p.relative_to(ROOT)), 'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
(OUT/'baseline_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
for p in list((ROOT/'source_downloads').glob('*.pdf'))+list((ROOT/'source_snapshots').glob('*.pdf')):
    d=fitz.open(p)
    text='\n'.join(f'\n===== PDF PAGE {i+1} =====\n'+page.get_text() for i,page in enumerate(d))
    (OUT/'extracted'/(p.stem+'.txt')).write_text(text,encoding='utf-8')
for p in (ROOT/'source_snapshots').glob('*.html'):
    (OUT/'extracted'/(p.stem+'.txt')).write_text(BeautifulSoup(p.read_text(encoding='utf-8'),'html.parser').get_text('\n',strip=True),encoding='utf-8')
for p in (ROOT/'source_downloads').glob('*.docx'):
    from docx import Document
    doc=Document(p)
    text='\n'.join(x.text for x in doc.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows)
    (OUT/'extracted'/(p.stem+'.txt')).write_text(text,encoding='utf-8')
sources=json.loads((ROOT/'sources.json').read_text(encoding='utf-8'))
print('SOURCES')
for k,s in sources.items(): print(k,s['url'],s['version_or_update_date'],s.get('local_file','NO_LOCAL'))
for name in ['verify_research.py','verify_negative_checks.py']:
    r=subprocess.run([sys.executable,'-X','utf8',str(ROOT/name)],capture_output=True,text=True,encoding='utf-8')
    (OUT/('run_'+name+'.stdout.txt')).write_text(r.stdout,encoding='utf-8')
    (OUT/('run_'+name+'.stderr.txt')).write_text(r.stderr,encoding='utf-8')
    report='verification_report.json' if name=='verify_research.py' else 'verification_negative_checks.json'
    (OUT/('independent_'+report)).write_bytes((ROOT/report).read_bytes())
    (ROOT/report).write_bytes((OUT/('original_'+report)).read_bytes())
    print(name,'EXIT',r.returncode)
rules=[yaml.safe_load(p.read_text(encoding='utf-8')) for p in sorted((ROOT/'rules').glob('*.yaml'))]
print('RULES',len(rules),'READY',sum(r['verification']['ready_for_technical_implementation'] for r in rules))
for r in rules:
    clean={k:v for k,v in r.items() if k!='sources'}
    (OUT/'extracted'/(r['rule_id']+'.json')).write_text(json.dumps(clean,ensure_ascii=False,indent=2),encoding='utf-8')
print('EXTRACTED',len(list((OUT/'extracted').iterdir())))
